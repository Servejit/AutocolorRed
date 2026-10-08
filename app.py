import streamlit as st
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from copy import copy
import re
import io


# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title="6thsense Vardaan Auto Color",
    layout="wide"
)

st.title("6thsense Vardaan Auto Color")


# ============================================================
# HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\u00a0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\u200c", "")
    text = text.replace("\u200d", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()


def get_number(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    text = text.replace(",", "").replace("%", "")

    try:
        return float(text)
    except:
        return None


def get_parentheses_number(value):
    if value is None:
        return None

    text = str(value)

    match = re.search(
        r"\(\s*([-+]?\d+(?:\.\d+)?)\s*\)",
        text
    )

    if match:
        try:
            return float(match.group(1))
        except:
            return None

    return None


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Excel file",
    type=["xlsx", "xlsm"]
)

if uploaded_file is None:
    st.info("Upload your master Excel file to continue.")
    st.stop()


# ============================================================
# LOAD ORIGINAL WORKBOOK
# ============================================================

file_bytes = uploaded_file.getvalue()

value_wb = openpyxl.load_workbook(
    io.BytesIO(file_bytes),
    data_only=True
)

format_wb = openpyxl.load_workbook(
    io.BytesIO(file_bytes),
    data_only=False
)

if "summary" not in value_wb.sheetnames:
    st.error("Sheet named 'summary' was not found.")
    st.stop()

value_ws = value_wb["summary"]
format_ws = format_wb["summary"]


# ============================================================
# CREATE OUTPUT WORKBOOK
# ============================================================

out_wb = Workbook()
out_ws = out_wb.active
out_ws.title = "AutoGreen"


# ============================================================
# COPY VALUES + FORMATTING
# ============================================================

# IMPORTANT: iter_rows() must be called as a method.
for row in value_ws.iter_rows():
    for value_cell in row:
        r = value_cell.row
        c = value_cell.column

        out_cell = out_ws.cell(
            row=r,
            column=c,
            value=value_cell.value
        )

        source_cell = format_ws.cell(row=r, column=c)

        if source_cell.has_style:
            out_cell.font = copy(source_cell.font)
            out_cell.fill = copy(source_cell.fill)
            out_cell.border = copy(source_cell.border)
            out_cell.alignment = copy(source_cell.alignment)
            out_cell.protection = copy(source_cell.protection)

        out_cell.number_format = source_cell.number_format


# ============================================================
# COPY COLUMN WIDTHS
# ============================================================

for key, dimension in format_ws.column_dimensions.items():
    out_ws.column_dimensions[key].width = dimension.width
    out_ws.column_dimensions[key].hidden = dimension.hidden


# ============================================================
# COPY ROW HEIGHTS
# ============================================================

for key, dimension in format_ws.row_dimensions.items():
    out_ws.row_dimensions[key].height = dimension.height
    out_ws.row_dimensions[key].hidden = dimension.hidden


# ============================================================
# COPY MERGED CELLS
# ============================================================

for merged_range in format_ws.merged_cells.ranges:
    out_ws.merge_cells(str(merged_range))


# ============================================================
# CREATE AUTOBLUE SHEET
# ============================================================

auto_blue_ws = out_wb.copy_worksheet(out_ws)
auto_blue_ws.title = "AutoBlue"


# ============================================================
# FIND HEADER ROW
# ============================================================

header_row = None

for row in range(1, out_ws.max_row + 1):
    if clean_text(out_ws.cell(row, 1).value) == "symbol":
        header_row = row
        break

if header_row is None:
    st.error("Could not find header row containing 'Symbol' in Column A.")
    st.stop()


# ============================================================
# AUTOBLUE SHEET
# ============================================================

auto_blue_header_row = None

for row in range(1, auto_blue_ws.max_row + 1):
    if clean_text(auto_blue_ws.cell(row, 1).value) == "symbol":
        auto_blue_header_row = row
        break

if auto_blue_header_row is not None:

    auto_blue_fill = PatternFill(fill_type="solid", fgColor="ADD8E6")
    auto_blue_new_stock_fill = PatternFill(fill_type="solid", fgColor="B4C6E7")
    auto_blue_condition_counts = {}
    auto_blue_new_match_rows = set()

    def make_auto_blue(row, col):
        auto_blue_ws.cell(row, col).fill = copy(auto_blue_fill)
        auto_blue_condition_counts[row] = (
            auto_blue_condition_counts.get(row, 0) + 1
        )

    def normalize_auto_blue_heading(value):
        text = clean_text(value)
        return re.sub(r"[^a-z0-9]+", "", text)

    def find_auto_blue_col(target_heading):
        target = normalize_auto_blue_heading(target_heading)
        for col in range(1, auto_blue_ws.max_column + 1):
            heading = normalize_auto_blue_heading(
                auto_blue_ws.cell(auto_blue_header_row, col).value
            )
            if heading == target:
                return col
        return None

    # 1. Sum I < -4
    col = find_auto_blue_col("sum i")
    if col:
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            number = get_number(auto_blue_ws.cell(row, col).value)
            if number is not None and number < -4:
                make_auto_blue(row, col)

    # 2. 16> C-B / Avg.4 — parentheses value < 0.50
    col = find_auto_blue_col("16> c-b / avg.4")
    if col:
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            value = get_parentheses_number(auto_blue_ws.cell(row, col).value)
            if value is not None and value < 0.50:
                make_auto_blue(row, col)

    # 3. 16< D-B / Avg.4 — parentheses < -1 and < number after 16<
    col = find_auto_blue_col("16< d-b / avg.4")
    if col:
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            cell_value = auto_blue_ws.cell(row, col).value
            if cell_value is None:
                continue

            text = str(cell_value)
            first_match = re.search(
                r"16<\s*([-+]?\d+(?:\.\d+)?)",
                text,
                re.IGNORECASE
            )
            parent_value = get_parentheses_number(cell_value)

            if first_match and parent_value is not None:
                try:
                    changing_value = float(first_match.group(1))
                except:
                    continue

                if parent_value < -1 and parent_value < changing_value:
                    make_auto_blue(row, col)

    # 4. Avg.4 O2H (set 2) and Avg.4 O2H (set 3)
    # Both values must be greater than the value inside ( )
    # in 16> C-B / Avg.4. If true, BOTH cells become blue.
    c_b_col = find_auto_blue_col("16> c-b / avg.4")
    set2_col = find_auto_blue_col("avg.4 o2h (set 2)")
    set3_col = find_auto_blue_col("avg.4 o2h (set 3)")

    if c_b_col and set2_col and set3_col:
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            c_b_parent = get_parentheses_number(
                auto_blue_ws.cell(row, c_b_col).value
            )
            set2_value = get_number(
                auto_blue_ws.cell(row, set2_col).value
            )
            set3_value = get_number(
                auto_blue_ws.cell(row, set3_col).value
            )

            if (
                c_b_parent is not None
                and set2_value is not None
                and set3_value is not None
                and set2_value > c_b_parent
                and set3_value > c_b_parent
            ):
                make_auto_blue(row, set2_col)
                make_auto_blue(row, set3_col)

    # 5. Sum O2H.10 — below average
    col = find_auto_blue_col("sum o2h.10")
    if col:
        values = []
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            number = get_number(auto_blue_ws.cell(row, col).value)
            if number is not None:
                values.append(number)

        if values:
            average_value = sum(values) / len(values)
            for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
                number = get_number(auto_blue_ws.cell(row, col).value)
                if number is not None and number < average_value:
                    make_auto_blue(row, col)

    # 6. Sum O2L.10 — below average
    col = find_auto_blue_col("sum o2l.10")
    if col:
        values = []
        for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
            number = get_number(auto_blue_ws.cell(row, col).value)
            if number is not None:
                values.append(number)

        if values:
            average_value = sum(values) / len(values)
            for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
                number = get_number(auto_blue_ws.cell(row, col).value)
                if number is not None and number < average_value:
                    make_auto_blue(row, col)

    # 7. AutoBlue gate
    dated_o2l_columns = []
    for col in range(1, auto_blue_ws.max_column + 1):
        heading = normalize_auto_blue_heading(
            auto_blue_ws.cell(auto_blue_header_row, col).value
        )
        if (
            "o2l" in heading
            and heading != normalize_auto_blue_heading("sum o2l.10")
        ):
            dated_o2l_columns.append(col)

    first_two_o2l = dated_o2l_columns[:2]
    negative_col = find_auto_blue_col('10 "-ve"')
    pct1_col = find_auto_blue_col("%chg.1")
    pct2_col = find_auto_blue_col("%chg.2")
    pct3_col = find_auto_blue_col("%chg.3")

    for row in range(auto_blue_header_row + 1, auto_blue_ws.max_row + 1):
        o2l_ok = False

        if len(first_two_o2l) >= 2:
            o2l_values = [
                get_number(auto_blue_ws.cell(row, col).value)
                for col in first_two_o2l
            ]
            o2l_ok = all(
                value is not None and value < -1.5
                for value in o2l_values
            )

        neg_ok = False
        if negative_col:
            value = auto_blue_ws.cell(row, negative_col).value
            if value is not None:
                # Correctly require a literal + after the first number.
                match = re.match(
                    r"^\s*(\d+)\s*\+",
                    str(value).strip()
                )
                neg_ok = bool(match and int(match.group(1)) > 3)

        pct_ok = False
        if pct1_col and pct2_col and pct3_col:
            pct_values = [
                get_number(auto_blue_ws.cell(row, col).value)
                for col in (pct1_col, pct2_col, pct3_col)
            ]
            pct_ok = all(
                value is not None and value < 0
                for value in pct_values
            )

        if o2l_ok and neg_ok and pct_ok:
            for col in first_two_o2l:
                make_auto_blue(row, col)

            make_auto_blue(row, negative_col)

            for col in (pct1_col, pct2_col, pct3_col):
                make_auto_blue(row, col)

            auto_blue_new_match_rows.add(row)

    for row in range(
        auto_blue_header_row + 1,
        auto_blue_ws.max_row + 1
    ):
        if row in auto_blue_new_match_rows:
            auto_blue_ws.cell(row, 1).fill = copy(auto_blue_new_stock_fill)
        else:
            auto_blue_ws.cell(row, 1).fill = PatternFill(fill_type=None)

    auto_blue_ws.auto_filter.ref = (
        f"A{auto_blue_header_row}:"
        f"{openpyxl.utils.get_column_letter(auto_blue_ws.max_column)}"
        f"{auto_blue_ws.max_row}"
    )
    auto_blue_ws.freeze_panes = f"A{auto_blue_header_row + 1}"


# ============================================================
# AUTOGREEN BLUE FILL
# ============================================================

blue_fill = PatternFill(fill_type="solid", fgColor="ADD8E6")
blue_cells = set()


def make_blue(row, col):
    out_ws.cell(row, col).fill = copy(blue_fill)
    blue_cells.add((row, col))
    out_ws.cell(row, 1).fill = copy(blue_fill)


# ============================================================
# AUTOGREEN HEADINGS
# ============================================================

headings = {}

for col in range(1, out_ws.max_column + 1):
    heading = clean_text(out_ws.cell(header_row, col).value)
    if heading:
        headings[heading] = col


# ============================================================
# RULE 1 — Sum I < -4
# ============================================================

sum_i_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "sum i":
        sum_i_col = col
        break

if sum_i_col:
    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, sum_i_col).value)
        if number is not None and number < -4:
            make_blue(row, sum_i_col)


# ============================================================
# RULE 2 — 16> C-B / Avg.4 parentheses < 0.50
# ============================================================

col_c = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "16> c-b / avg.4":
        col_c = col
        break

if col_c:
    for row in range(header_row + 1, out_ws.max_row + 1):
        value = get_parentheses_number(
            out_ws.cell(row, col_c).value
        )

        if value is not None and value < 0.50:
            make_blue(row, col_c)


# ============================================================
# RULE 3 — 16< D-B / Avg.4
# ============================================================

col_d = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "16< d-b / avg.4":
        col_d = col
        break

if col_d:
    for row in range(header_row + 1, out_ws.max_row + 1):
        cell_value = out_ws.cell(row, col_d).value

        if cell_value is None:
            continue

        text = str(cell_value)

        first_match = re.search(
            r"16<\s*([-+]?\d+(?:\.\d+)?)",
            text,
            re.IGNORECASE
        )

        parent_value = get_parentheses_number(cell_value)

        if first_match and parent_value is not None:
            try:
                changing_value = float(first_match.group(1))
            except:
                continue

            if parent_value < -1 and parent_value < changing_value:
                make_blue(row, col_d)


# ============================================================
# RULE 4 — NEW SET 2 / SET 3 CONDITION
#
# Avg.4 O2H (set 2) > value inside ( )
# of 16> C-B / Avg.4
#
# AND
#
# Avg.4 O2H (set 3) > value inside ( )
# of 16> C-B / Avg.4
#
# When BOTH are true, BOTH cells become blue.
# ============================================================

set2_col = None
set3_col = None

for col in range(1, out_ws.max_column + 1):
    heading = clean_text(out_ws.cell(header_row, col).value)

    if heading == "avg.4 o2h (set 2)":
        set2_col = col

    elif heading == "avg.4 o2h (set 3)":
        set3_col = col

if col_c and set2_col and set3_col:
    for row in range(header_row + 1, out_ws.max_row + 1):

        c_b_parent = get_parentheses_number(
            out_ws.cell(row, col_c).value
        )

        set2_value = get_number(
            out_ws.cell(row, set2_col).value
        )

        set3_value = get_number(
            out_ws.cell(row, set3_col).value
        )

        if (
            c_b_parent is not None
            and set2_value is not None
            and set3_value is not None
            and set2_value > c_b_parent
            and set3_value > c_b_parent
        ):
            make_blue(row, set2_col)
            make_blue(row, set3_col)


# ============================================================
# RULE 5 — Sum O2H.10 below average
# ============================================================

sum_o2h_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "sum o2h.10":
        sum_o2h_col = col
        break

if sum_o2h_col:
    values = []

    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, sum_o2h_col).value)

        if number is not None:
            values.append(number)

    if values:
        average_value = sum(values) / len(values)

        for row in range(header_row + 1, out_ws.max_row + 1):
            number = get_number(
                out_ws.cell(row, sum_o2h_col).value
            )

            if number is not None and number < average_value:
                make_blue(row, sum_o2h_col)


# ============================================================
# RULE 6 — Sum O2L.10 below average
# ============================================================

sum_o2l_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "sum o2l.10":
        sum_o2l_col = col
        break

if sum_o2l_col:
    values = []

    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, sum_o2l_col).value)

        if number is not None:
            values.append(number)

    if values:
        average_value = sum(values) / len(values)

        for row in range(header_row + 1, out_ws.max_row + 1):
            number = get_number(
                out_ws.cell(row, sum_o2l_col).value
            )

            if number is not None and number < average_value:
                make_blue(row, sum_o2l_col)


# ============================================================
# RULE 7 — FIRST TWO DATED O2L COLUMNS < -1
# ============================================================

dated_o2l_columns = []

for col in range(1, out_ws.max_column + 1):
    heading = clean_text(out_ws.cell(header_row, col).value)

    if "o2l" in heading and heading != "sum o2l.10":
        dated_o2l_columns.append(col)

for col in dated_o2l_columns[:2]:
    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, col).value)

        if number is not None and number < -1:
            make_blue(row, col)


# ============================================================
# RULE 8 — 10 "-ve", first number before + > 1
# ============================================================

negative_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == '10 "-ve"':
        negative_col = col
        break

if negative_col:
    for row in range(header_row + 1, out_ws.max_row + 1):
        value = out_ws.cell(row, negative_col).value

        if value is None:
            continue

        match = re.match(
            r"^\s*(\d+)\s*\+",
            str(value).strip()
        )

        if match:
            try:
                first_number = int(match.group(1))
            except:
                continue

            if first_number > 1:
                make_blue(row, negative_col)


# ============================================================
# RULE 9 — %Chg.1 < 0
# ============================================================

pct1_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "%chg.1":
        pct1_col = col
        break

if pct1_col:
    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, pct1_col).value)

        if number is not None and number < 0:
            make_blue(row, pct1_col)


# ============================================================
# RULE 10 — %Chg.2 < 0
# ============================================================

pct2_col = None

for col in range(1, out_ws.max_column + 1):
    if clean_text(out_ws.cell(header_row, col).value) == "%chg.2":
        pct2_col = col
        break

if pct2_col:
    for row in range(header_row + 1, out_ws.max_row + 1):
        number = get_number(out_ws.cell(row, pct2_col).value)

        if number is not None and number < 0:
            make_blue(row, pct2_col)


# ============================================================
# COUNT BLUE CELLS PER ROW
# ============================================================

row_blue_counts = {}

for row in range(header_row + 1, out_ws.max_row + 1):
    count = sum(
        1
        for r, c in blue_cells
        if r == row and c != 1
    )

    row_blue_counts[row] = count


# ============================================================
# THREE GREEN LEVELS
# ============================================================

dark_green_fill = PatternFill(
    fill_type="solid",
    fgColor="548235"
)

green_fill = PatternFill(
    fill_type="solid",
    fgColor="70AD47"
)

light_green_fill = PatternFill(
    fill_type="solid",
    fgColor="90EE90"
)

eligible_counts = sorted(
    {
        count
        for count in row_blue_counts.values()
        if count >= 8
    },
    reverse=True
)

green_counts = eligible_counts[:3]

for row, count in row_blue_counts.items():

    if count < 8:
        continue

    if len(green_counts) >= 1 and count == green_counts[0]:
        out_ws.cell(row, 1).fill = copy(dark_green_fill)

    elif len(green_counts) >= 2 and count == green_counts[1]:
        out_ws.cell(row, 1).fill = copy(green_fill)

    elif len(green_counts) >= 3 and count == green_counts[2]:
        out_ws.cell(row, 1).fill = copy(light_green_fill)


# ============================================================
# AUTOFILTER + FREEZE
# ============================================================

out_ws.auto_filter.ref = (
    f"A{header_row}:"
    f"{openpyxl.utils.get_column_letter(out_ws.max_column)}"
    f"{out_ws.max_row}"
)

out_ws.freeze_panes = f"A{header_row + 1}"


# ============================================================
# SAVE
# ============================================================

output_buffer = io.BytesIO()

out_wb.save(output_buffer)

output_buffer.seek(0)

st.success("Excel file processed successfully.")

st.download_button(
    label="Download 6thsenseVardaanAutocolor.xlsx",
    data=output_buffer.getvalue(),
    file_name="6thsenseVardaanAutocolor.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
